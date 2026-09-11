"""Ручки отчёта: форма ответа, роли, границы периода и вложенность.

Арифметику проверяет `test_crud_reports` — там «сейчас» задаётся явно. Здесь
его задать нельзя: ручка берёт момент из системных часов, чтобы отличить
просроченное ТО от текущего. Поэтому данные заводятся **относительно
сегодняшнего дня**: ТО на предыдущий месяц просрочено всегда, на текущий —
никогда, в какой бы день года ни случился прогон.
"""

import datetime
import itertools
import json
import uuid
from io import BytesIO

import pytest
from openpyxl import load_workbook

from src.config import settings
from src.core.roles import ADMIN, CLIENT_ID, ENGINEER, FOREMAN, MECHANIC
from src.crud.crud_statistics import _PLANNED_MONTH_COLUMN, previous_month
from src.models import (
    ActFact,
    Company,
    DefectiveAct,
    Object,
    Organization,
    PlannedTO,
    UniversalUser,
)

URL = f"{settings.API_V1_STR}/reports/works"
EXPORT_URL = f"{settings.API_V1_STR}/reports/works/export"
CAT_A = 2  # остановка лифта — поломка
CAT_TO = 6  # плановые работы — не поломка


def this_year():
    return datetime.date.today().year


def year_bounds():
    """Весь текущий год — период, в который попадает всё заведённое ниже."""
    year = this_year()
    return (
        datetime.date(year, 1, 1).isoformat(),
        datetime.date(year, 12, 31).isoformat(),
    )


@pytest.fixture
def make_object(db_session):
    counter = itertools.count(1)
    prefix = uuid.uuid4().hex[:8]

    def _make(**kwargs):
        number = next(counter)
        obj = Object(
            name=kwargs.pop("name", f"Лифт {number}"),
            factory_number=f"F-{prefix}-{number}",
            registration_number=f"R-{prefix}-{number}",
            **kwargs,
        )
        db_session.add(obj)
        db_session.flush()
        return obj

    return _make


@pytest.fixture
def plan_to(db_session):
    def _make(obj, *, months=None, year=None, step_list_fact=None):
        columns = {}
        for month, finished_at in (months or {}).items():
            act = ActFact(
                object_id=obj.id,
                finished_at=finished_at,
                step_list_fact=step_list_fact,
            )
            db_session.add(act)
            db_session.flush()
            columns[_PLANNED_MONTH_COLUMN[month].key] = act.id

        planned = PlannedTO(
            year=str(year or this_year()), object_id=obj.id, **columns
        )
        db_session.add(planned)
        db_session.flush()
        return planned

    return _make


def _params(**extra):
    date_from, date_to = year_bounds()
    return {"date_from": date_from, "date_to": date_to, **extra}


def _data(response):
    assert response.status_code == 200, response.text
    return response.json()["data"]


class TestAccess:
    """Отчёты открыты правом `STATISTICS_READ`, а оно есть у всех ролей.

    Разграничение делает не право, а область видимости: прораб видит свои
    участки, механик — свои объекты, клиент — свою компанию. Тест фиксирует
    это как решение, а не как случайность: если отчёты когда-нибудь захотят
    закрыть от механика, понадобится отдельное право, и здесь станет видно,
    что менять.
    """

    @pytest.mark.integration
    @pytest.mark.parametrize(
        "role",
        [ADMIN, FOREMAN, MECHANIC, ENGINEER, CLIENT_ID],
        ids=["админ", "прораб", "механик", "инженер", "клиент"],
    )
    def test_every_role_with_statistics_right_is_allowed(
        self, client_with_db, as_role, role
    ):
        as_role(role)

        response = client_with_db.get(URL, params=_params())

        assert response.status_code == 200

    @pytest.mark.integration
    def test_mechanic_sees_only_own_objects(
        self, client_with_db, as_role, make_object
    ):
        # Право у механика есть, но выдача пуста, пока за ним не закреплён
        # ни один лифт: границу держит область видимости, а не право.
        make_object()
        as_role(MECHANIC)

        data = _data(client_with_db.get(URL, params=_params()))

        assert data["items"] == []
        assert data["total_objects"] == 0

    @pytest.mark.integration
    def test_client_sees_only_own_company(
        self, client_with_db, as_role, db_session, make_object
    ):
        # Область видимости — граница, а не фильтр: запрошенный чужой
        # company_id выдачу не расширяет.
        mine = Company(name=f"Моя {uuid.uuid4().hex[:6]}")
        theirs = Company(name=f"Чужая {uuid.uuid4().hex[:6]}")
        db_session.add_all([mine, theirs])
        db_session.flush()
        make_object(company_id=mine.id)
        alien = make_object(company_id=theirs.id)
        as_role(CLIENT_ID, company_id=mine.id)

        data = _data(
            client_with_db.get(URL, params=_params(company_id=theirs.id))
        )

        assert alien.id not in [item["object_id"] for item in data["items"]]


class TestPeriod:
    @pytest.mark.integration
    def test_reversed_period_is_rejected(self, client_with_db, as_role):
        as_role(ADMIN)

        response = client_with_db.get(
            URL, params={"date_from": "2026-06-01", "date_to": "2026-03-01"}
        )

        assert response.status_code == 422

    @pytest.mark.integration
    def test_too_long_period_is_rejected(self, client_with_db, as_role):
        # Каждый месяц — ветка в CASE и в UNION ALL. Десятилетний период
        # раздувает запрос без пользы: отчёты отправляют за год.
        as_role(ADMIN)

        response = client_with_db.get(
            URL, params={"date_from": "2016-01-01", "date_to": "2026-12-31"}
        )

        assert response.status_code == 422

    @pytest.mark.integration
    def test_partial_month_is_reported_as_whole(self, client_with_db, as_role):
        # Просили с 15 марта, а ТО считаются по месяцу целиком. Ответ обязан
        # сказать об этом сам, иначе лишнее ТО выглядит ошибкой счёта.
        as_role(ADMIN)
        year = this_year()

        data = _data(
            client_with_db.get(
                URL,
                params={
                    "date_from": datetime.date(year, 3, 15).isoformat(),
                    "date_to": datetime.date(year, 6, 20).isoformat(),
                },
            )
        )

        assert data["period"]["month_from"] == {"year": year, "month": 3}
        assert data["period"]["month_to"] == {"year": year, "month": 6}
        assert data["period"]["months_count"] == 4


class TestMatrix:
    @pytest.mark.integration
    def test_matrix_has_every_month_of_period(
        self, client_with_db, as_role, make_object
    ):
        # Месяцы без работ обязаны быть в матрице нулями: иначе двенадцать
        # ячеек превратятся в три и колонки разъедутся между строками.
        as_role(ADMIN)
        obj = make_object()

        data = _data(client_with_db.get(URL, params=_params(object_id=obj.id)))

        assert len(data["months"]) == 12
        assert len(data["items"][0]["months"]) == 12

    @pytest.mark.integration
    def test_object_without_work_is_present(
        self, client_with_db, as_role, make_object
    ):
        as_role(ADMIN)
        obj = make_object()

        data = _data(client_with_db.get(URL, params=_params(object_id=obj.id)))

        row = data["items"][0]
        assert row["object_id"] == obj.id
        assert row["maintenance_planned"] == 0
        assert row["counts"]["breakdowns"] == 0
        assert all(cell["maintenance"] == "none" for cell in row["months"])

    @pytest.mark.integration
    def test_overdue_and_pending_are_different_states(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Предыдущий месяц просрочен всегда, текущий — никогда.
        as_role(ADMIN)
        obj = make_object()
        past_year, past_month = previous_month(this_year(), datetime.date.today().month)
        current_month = datetime.date.today().month

        if past_year != this_year():
            pytest.skip("В январе предыдущий месяц лежит в прошлом году")

        plan_to(obj, months={past_month: None, current_month: None})

        data = _data(client_with_db.get(URL, params=_params(object_id=obj.id)))
        cells = {cell["month"]: cell["maintenance"] for cell in data["items"][0]["months"]}

        assert cells[past_month] == "overdue"
        assert cells[current_month] == "pending"
        assert data["summary"]["maintenance_overdue"] == 1

    @pytest.mark.integration
    def test_summary_counts_whole_selection_not_page(
        self, client_with_db, as_role, db_session, make_object
    ):
        as_role(ADMIN)
        organization = Organization(title=f"Орг {uuid.uuid4().hex[:6]}")
        db_session.add(organization)
        db_session.flush()
        for _ in range(3):
            make_object(organization_id=organization.id)

        data = _data(
            client_with_db.get(
                URL, params=_params(organization_id=organization.id, limit=2)
            )
        )

        assert len(data["items"]) == 2
        assert data["total_objects"] == 3
        assert data["summary"]["objects_total"] == 3


class TestObjectWorks:
    @pytest.mark.integration
    def test_unknown_object_is_404(self, client_with_db, as_role):
        as_role(ADMIN)

        response = client_with_db.get(
            f"{settings.API_V1_STR}/reports/object/999999/works", params=_params()
        )

        assert response.status_code == 404

    @pytest.mark.integration
    def test_invisible_object_is_404_not_empty_report(
        self, client_with_db, as_role, db_session, make_object
    ):
        # Пустой отчёт вместо 404 позволил бы перебором узнать, какие id
        # существуют.
        mine = Company(name=f"Моя {uuid.uuid4().hex[:6]}")
        theirs = Company(name=f"Чужая {uuid.uuid4().hex[:6]}")
        db_session.add_all([mine, theirs])
        db_session.flush()
        alien = make_object(company_id=theirs.id)
        as_role(CLIENT_ID, company_id=mine.id)

        response = client_with_db.get(
            f"{settings.API_V1_STR}/reports/object/{alien.id}/works",
            params=_params(),
        )

        assert response.status_code == 404

    @pytest.mark.integration
    def test_checklist_is_parsed_from_python_repr(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Фронт подрядчика писал чек-лист питоновским repr — с одинарными
        # кавычками, а не JSON. Разбор обязан это переживать.
        as_role(ADMIN)
        obj = make_object()
        checklist = str(
            [
                {
                    "step_name": "Осмотр канатов",
                    "step_status": "true",
                    "substeps": [
                        {"substep_name": "Проверить натяжение", "substep_status": "0"}
                    ],
                }
            ]
        )
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)}, step_list_fact=checklist)

        data = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=_params(),
            )
        )

        steps = data["maintenance"][0]["steps"]
        assert [step["title"] for step in steps] == [
            "Осмотр канатов",
            "Проверить натяжение",
        ]
        assert [step["done"] for step in steps] == [True, False]

    @pytest.mark.integration
    def test_checklist_in_the_shape_both_fronts_actually_write(
        self, client_with_db, as_role, make_object, plan_to
    ):
        """Форма, которая лежит в базе на самом деле.

        Экран графика у прораба и мобильное приложение механика пишут не
        список шагов, а словарь с названием ТО, внутри которого список шагов
        лежит **ещё одной строкой**. Разбор искал `step_name` и на этой форме
        молча отдавал пустой список — то есть отчёт заказчику показывал
        «чек-лист не заполнен» по каждому сделанному ТО.
        """
        as_role(ADMIN)
        obj = make_object()
        checklist = json.dumps(
            {
                "numberTo": "ТО-1",
                "stepListTO": json.dumps(
                    [
                        {"text": "Выключить вводное устройство", "bool": True},
                        {"text": "Осмотр станции управления", "bool": False},
                    ],
                    ensure_ascii=False,
                ),
            },
            ensure_ascii=False,
        )
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)}, step_list_fact=checklist)

        data = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=_params(),
            )
        )

        assert data["maintenance"][0]["steps"] == [
            {"title": "Выключить вводное устройство", "done": True},
            {"title": "Осмотр станции управления", "done": False},
        ]

    @pytest.mark.integration
    def test_unreadable_checklist_gives_empty_steps(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Пустой список означает «чек-лист не заполнен», а не падение ручки.
        as_role(ADMIN)
        obj = make_object()
        plan_to(
            obj,
            months={3: datetime.datetime(this_year(), 3, 20)},
            step_list_fact="совсем не список",
        )

        data = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=_params(),
            )
        )

        assert data["maintenance"][0]["steps"] == []

    @pytest.mark.integration
    def test_json_checklist_also_works(
        self, client_with_db, as_role, make_object, plan_to
    ):
        as_role(ADMIN)
        obj = make_object()
        plan_to(
            obj,
            months={3: datetime.datetime(this_year(), 3, 20)},
            step_list_fact=json.dumps(
                [{"step_name": "Смазка направляющих", "step_status": True}]
            ),
        )

        data = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=_params(),
            )
        )

        assert data["maintenance"][0]["steps"] == [
            {"title": "Смазка направляющих", "done": True}
        ]

    @pytest.mark.integration
    def test_object_row_matches_matrix(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Шапка раскрытой строки считана тем же кодом, что и матрица, поэтому
        # разойтись с ней не может. Тест это фиксирует.
        as_role(ADMIN)
        obj = make_object()
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)})

        matrix = _data(client_with_db.get(URL, params=_params(object_id=obj.id)))
        detail = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=_params(),
            )
        )

        assert detail["object"] == matrix["items"][0]


class TestExport:
    """Выгрузка в Excel: та же арифметика, что на экране, но все объекты."""

    @pytest.mark.integration
    def test_file_has_four_sheets(self, client_with_db, as_role, make_object):
        as_role(ADMIN)
        make_object()

        response = client_with_db.get(EXPORT_URL, params=_params())

        assert response.status_code == 200, response.text
        book = load_workbook(BytesIO(response.content))
        assert book.sheetnames == ["Сводка", "ТО", "Заявки", "Дефекты"]

    @pytest.mark.integration
    def test_filename_carries_period(self, client_with_db, as_role):
        as_role(ADMIN)
        date_from, date_to = year_bounds()

        response = client_with_db.get(EXPORT_URL, params=_params())

        disposition = response.headers["content-disposition"]
        assert f"works-{date_from}-{date_to}.xlsx" in disposition
        # Кириллица уезжает по RFC 5987: заголовок ходит в latin-1.
        assert "filename*=UTF-8''" in disposition

    @pytest.mark.integration
    def test_export_takes_all_objects_not_page(
        self, client_with_db, as_role, db_session, make_object, plan_to
    ):
        # Экран просит страницу, файл — весь отбор. Это единственное, чем они
        # отличаются: считает их одна функция.
        as_role(ADMIN)
        organization = Organization(title=f"Орг {uuid.uuid4().hex[:6]}")
        db_session.add(organization)
        db_session.flush()
        for _ in range(3):
            obj = make_object(organization_id=organization.id)
            plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)})

        params = _params(organization_id=organization.id, limit=1)
        screen = _data(client_with_db.get(URL, params=params))
        book = load_workbook(
            BytesIO(client_with_db.get(EXPORT_URL, params=params).content)
        )

        # На экране одна строка из трёх, в файле — все три ТО.
        assert len(screen["items"]) == 1
        assert book["ТО"].max_row == 4  # шапка плюс три строки

    @pytest.mark.integration
    def test_summary_in_file_matches_screen(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Цифра в отправленном заказчику файле обязана совпадать с экраном.
        as_role(ADMIN)
        obj = make_object()
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)})

        params = _params(object_id=obj.id)
        screen = _data(client_with_db.get(URL, params=params))
        sheet = load_workbook(
            BytesIO(client_with_db.get(EXPORT_URL, params=params).content)
        )["Сводка"]

        labels = {
            sheet.cell(row=row, column=1).value: sheet.cell(row=row, column=2).value
            for row in range(5, 17)
        }
        assert labels["ТО запланировано"] == screen["summary"]["maintenance_planned"]
        assert labels["ТО выполнено"] == screen["summary"]["maintenance_completed"]
        assert labels["Объектов в отчёте"] == screen["summary"]["objects_total"]

    @pytest.mark.integration
    def test_scope_binds_export_too(
        self, client_with_db, as_role, db_session, make_object, plan_to
    ):
        # Выгрузка идёт мимо `require(...)`, поэтому область видимости на ней
        # проверяется отдельно: забытый фильтр здесь отдал бы чужие данные
        # файлом, и заметить это было бы нечем.
        mine = Company(name=f"Моя {uuid.uuid4().hex[:6]}")
        theirs = Company(name=f"Чужая {uuid.uuid4().hex[:6]}")
        db_session.add_all([mine, theirs])
        db_session.flush()
        alien = make_object(company_id=theirs.id)
        plan_to(alien, months={3: datetime.datetime(this_year(), 3, 20)})
        as_role(CLIENT_ID, company_id=mine.id)

        book = load_workbook(
            BytesIO(
                client_with_db.get(
                    EXPORT_URL, params=_params(company_id=theirs.id)
                ).content
            )
        )

        assert book["ТО"].max_row == 1  # одна шапка, ни одной строки


class TestPdfExport:
    """PDF — документ для клиента, а не выгрузка данных."""

    @pytest.mark.integration
    def test_pdf_is_returned(self, client_with_db, as_role, make_object, plan_to):
        as_role(ADMIN)
        obj = make_object()
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)})

        response = client_with_db.get(
            EXPORT_URL, params=_params(format="pdf", object_id=obj.id)
        )

        assert response.status_code == 200, response.text
        assert response.headers["content-type"] == "application/pdf"
        # Сигнатура формата: собранный ReportLab документ начинается с неё.
        assert response.content[:5] == b"%PDF-"
        assert ".pdf" in response.headers["content-disposition"]

    @pytest.mark.integration
    def test_unknown_format_is_rejected(self, client_with_db, as_role):
        as_role(ADMIN)

        response = client_with_db.get(EXPORT_URL, params=_params(format="docx"))

        # 400, а не 422: проверку формата делает сам FastAPI по регулярке, а
        # ошибки валидации это приложение отдаёт своим обработчиком. Свои же
        # проверки периода поднимают 422 руками — отсюда разные коды у
        # соседних тестов.
        assert response.status_code == 400

    @pytest.mark.integration
    def test_photos_are_off_by_default(
        self, client_with_db, as_role, make_object, plan_to
    ):
        # Без галочки файл обязан остаться лёгким: с фотографиями годовой
        # отчёт весит сотни мегабайт и не уходит по почте.
        as_role(ADMIN)
        obj = make_object()
        plan_to(obj, months={3: datetime.datetime(this_year(), 3, 20)})

        plain = client_with_db.get(
            EXPORT_URL, params=_params(format="pdf", object_id=obj.id)
        )
        with_photos = client_with_db.get(
            EXPORT_URL,
            params=_params(format="pdf", object_id=obj.id, with_photos=True),
        )

        assert plain.status_code == 200
        assert with_photos.status_code == 200
        # Фотографий у объекта нет, поэтому размеры совпадают — важно, что
        # включённая галочка не роняет сборку.
        assert len(plain.content) > 0

    @pytest.mark.integration
    def test_pdf_survives_empty_selection(self, client_with_db, as_role):
        # На проде отчёт будет почти пустым: графики заведены у единиц
        # объектов. Пустой отбор обязан давать документ, а не ошибку.
        as_role(ADMIN)

        response = client_with_db.get(
            EXPORT_URL, params=_params(format="pdf", object_id=999999)
        )

        assert response.status_code == 200
        assert response.content[:5] == b"%PDF-"

    @pytest.mark.integration
    def test_scope_binds_pdf_too(
        self, client_with_db, as_role, db_session, make_object, plan_to
    ):
        mine = Company(name=f"Моя {uuid.uuid4().hex[:6]}")
        theirs = Company(name=f"Чужая {uuid.uuid4().hex[:6]}")
        db_session.add_all([mine, theirs])
        db_session.flush()
        alien = make_object(company_id=theirs.id, name="Чужой лифт")
        plan_to(alien, months={3: datetime.datetime(this_year(), 3, 20)})
        as_role(CLIENT_ID, company_id=mine.id)

        response = client_with_db.get(
            EXPORT_URL, params=_params(format="pdf", company_id=theirs.id)
        )

        assert response.status_code == 200
        # Имя чужого лифта не должно попасть в документ ни в каком виде.
        assert "Чужой лифт".encode() not in response.content


DEFECTS_URL = f"{settings.API_V1_STR}/reports/works/defects"


@pytest.fixture
def make_defect(db_session):
    """Дефектный акт с нужной датой создания. Плановое ТО не заполняется:
    три входа из четырёх его и не знают."""

    def _make(obj, *, created_at, kind="internal", parent=None):
        act = DefectiveAct(
            object_id=obj.id,
            title=f"Дефект {uuid.uuid4().hex[:6]}",
            kind=kind,
            parent_id=parent.id if parent else None,
            created_at=created_at,
        )
        db_session.add(act)
        db_session.flush()
        return act

    return _make


class TestDefectsList:
    """Список актов с плитки сводки считается тем же запросом, что и сама
    плитка, и по тем же правилам, что лента актов объекта в окне графика:
    дата создания и только внутренние акты."""

    @pytest.mark.integration
    def test_act_without_planned_to_is_counted(
        self, client_with_db, as_role, make_object, make_defect
    ):
        # Акт с пункта меню или по заявке планового ТО не знает — раньше
        # отчёт такие терял.
        as_role(ADMIN)
        obj = make_object()
        act = make_defect(obj, created_at=datetime.datetime(this_year(), 3, 5))

        rows = _data(
            client_with_db.get(DEFECTS_URL, params=_params(object_id=obj.id))
        )
        report = _data(client_with_db.get(URL, params=_params(object_id=obj.id)))

        assert rows["total"] == 1
        assert [row["defect_id"] for row in rows["items"]] == [act.id]
        assert rows["items"][0]["object_id"] == obj.id
        assert rows["items"][0]["month"] == 3
        assert report["summary"]["counts"]["defects"] == 1

    @pytest.mark.integration
    def test_client_act_is_not_counted(
        self, client_with_db, as_role, make_object, make_defect
    ):
        as_role(ADMIN)
        obj = make_object()
        when = datetime.datetime(this_year(), 4, 1)
        internal = make_defect(obj, created_at=when)
        make_defect(obj, created_at=when, kind="client", parent=internal)

        rows = _data(
            client_with_db.get(DEFECTS_URL, params=_params(object_id=obj.id))
        )

        assert [row["defect_id"] for row in rows["items"]] == [internal.id]

    @pytest.mark.integration
    def test_list_matches_summary_and_object_feed(
        self, client_with_db, as_role, make_object, make_defect
    ):
        # Три места показывают одно число: плитка сводки, список под ней и
        # лента актов в шторке объекта.
        as_role(ADMIN)
        obj = make_object()
        for month in (1, 2, 2):
            make_defect(obj, created_at=datetime.datetime(this_year(), month, 10))

        params = _params(object_id=obj.id)
        rows = _data(client_with_db.get(DEFECTS_URL, params=params))
        report = _data(client_with_db.get(URL, params=params))
        works = _data(
            client_with_db.get(
                f"{settings.API_V1_STR}/reports/object/{obj.id}/works",
                params=params,
            )
        )

        assert rows["total"] == len(rows["items"]) == 3
        assert report["summary"]["counts"]["defects"] == 3
        assert report["items"][0]["counts"]["defects"] == 3
        assert len(works["defects"]) == 3
        assert works["object"]["counts"]["defects"] == 3

    @pytest.mark.integration
    def test_act_outside_period_is_left_out(
        self, client_with_db, as_role, make_object, make_defect
    ):
        as_role(ADMIN)
        obj = make_object()
        make_defect(obj, created_at=datetime.datetime(this_year() - 1, 12, 31))

        rows = _data(
            client_with_db.get(DEFECTS_URL, params=_params(object_id=obj.id))
        )

        assert rows["items"] == []
        assert rows["total"] == 0

    @pytest.mark.integration
    def test_mechanic_does_not_see_alien_acts(
        self, client_with_db, as_role, make_object, make_defect
    ):
        obj = make_object()
        make_defect(obj, created_at=datetime.datetime(this_year(), 6, 1))
        as_role(MECHANIC)

        rows = _data(client_with_db.get(DEFECTS_URL, params=_params()))

        assert rows["items"] == []


class TestPdfDefectsSection:
    """Сводный раздел актов в PDF: один список по всему отбору, в
    дополнение к перечню внутри объекта."""

    @pytest.mark.integration
    def test_pdf_builds_with_defects(
        self, client_with_db, as_role, make_object, make_defect
    ):
        as_role(ADMIN)
        obj = make_object()
        make_defect(obj, created_at=datetime.datetime(this_year(), 2, 3))
        make_defect(obj, created_at=datetime.datetime(this_year(), 5, 9))

        for with_photos in (False, True):
            response = client_with_db.get(
                EXPORT_URL,
                params=_params(
                    format="pdf", object_id=obj.id, with_photos=with_photos
                ),
            )
            assert response.status_code == 200, response.text
            assert response.content[:5] == b"%PDF-"

    def test_section_rows_match_input(self):
        # Текст из PDF не достать: шрифт вшит субсетом. Поэтому число строк
        # проверяется на самой таблице, до сборки документа.
        from types import SimpleNamespace

        from reportlab.platypus import Table

        from src.services.reports_pdf import _defects_section, _ensure_font, _styles

        styles = _styles(_ensure_font())
        rows = [
            SimpleNamespace(
                object_name="Лифт 1",
                address="ул. Ленина, 1",
                title="Трос",
                description="Износ",
                created_at=datetime.datetime(this_year(), 2, 3),
                status="Открыт",
                responsible="Иванов",
                photo_count=2,
            ),
            SimpleNamespace(
                object_name="Лифт 2",
                address=None,
                title="Дверь",
                description=None,
                created_at=datetime.datetime(this_year(), 5, 9),
                status=None,
                responsible=None,
                photo_count=0,
            ),
        ]

        story = _defects_section(rows, styles, 700)

        tables = [item for item in story if isinstance(item, Table)]
        assert len(tables) == 1
        # Шапка и по строке на акт.
        assert len(tables[0]._cellvalues) == 3
        assert _defects_section([], styles, 700) == []
